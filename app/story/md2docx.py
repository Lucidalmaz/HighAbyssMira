"""Story-Bibel (Markdown) -> Word-Datei für den Autor. Überschriften, Listen, Fettdruck; '✏️ DEINE IDEEN:'-Stellen als gelb markierte Schreibfelder."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_COLOR_INDEX
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, dst = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Georgia'; st.font.size = Pt(11)
for s in doc.sections: s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
for lvl, size, col in [(1, 22, (0x4a, 0x10, 0x10)), (2, 16, (0x5a, 0x1a, 0x1a)), (3, 13, (0x33, 0x33, 0x33))]:
    h = doc.styles[f'Heading {lvl}']; h.font.name = 'Georgia'; h.font.size = Pt(size); h.font.color.rgb = RGBColor(*col)

def runs(par, text):
    # **fett**, *kursiv*, `code`
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text):
        if not part: continue
        if part.startswith('**'): r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and len(part) > 2: r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith('`'): r = par.add_run(part[1:-1]); r.font.name = 'Consolas'
        else: r = par.add_run(part)

def idea_box(label):
    p = doc.add_paragraph(); r = p.add_run('✏️ ' + label.strip()); r.bold = True; r.font.color.rgb = RGBColor(0x8a, 0x5a, 0x00)
    for _ in range(3):
        q = doc.add_paragraph(); rr = q.add_run(' ' * 120); rr.font.highlight_color = WD_COLOR_INDEX.YELLOW

lines = open(src, encoding='utf-8').read().splitlines()
first = True
for ln in lines:
    s = ln.rstrip()
    if not s.strip(): continue
    if s.startswith('# '):
        p = doc.add_heading(s[2:].strip(), 1 if not first else 0); first = False
        if p.style.name == 'Title': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    m = re.match(r'^(#{2,4})\s+(.*)', s)
    if m: doc.add_heading(m.group(2).strip(), min(3, len(m.group(1)))); continue
    if '✏️' in s and 'DEINE IDEEN' in s.upper():
        idea_box(re.sub(r'^[-*>\s]*', '', s)); continue
    m = re.match(r'^(\s*)[-*]\s+(.*)', s)
    if m:
        lvl = len(m.group(1)) // 2; p = doc.add_paragraph(style='List Bullet' if lvl == 0 else 'List Bullet 2'); runs(p, m.group(2)); continue
    m = re.match(r'^(\s*)\d+[.)]\s+(.*)', s)
    if m: p = doc.add_paragraph(style='List Number'); runs(p, m.group(2)); continue
    if s.startswith('>'): p = doc.add_paragraph(); r = p.add_run(s.lstrip('> ')); r.italic = True; continue
    if re.match(r'^-{3,}$', s): doc.add_paragraph('—' * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER; continue
    p = doc.add_paragraph(); runs(p, s)
doc.save(dst); print('gespeichert:', dst)
